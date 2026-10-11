"""Leitura das planilhas. Só lê — nunca grava nas planilhas do Farley."""
import os
import warnings

import openpyxl

from .util import nao, txt

warnings.filterwarnings('ignore', module='openpyxl')


def _abrir(caminho):
    return openpyxl.load_workbook(caminho, data_only=True, read_only=False)


def ler_selecao(caminho):
    """RNQ-selecao.xlsx -> {'pessoas': [...], 'comunidades': [...], 'ligacoes': [(a, b), ...]}"""
    wb = _abrir(caminho)
    comunidades = [txt(c.value) for c in wb['Comunidades']['A'][1:] if txt(c.value)]
    ws = wb['Pessoas']
    col = {txt(c.value): i for i, c in enumerate(ws[1])}
    obrig = ['Incluir?', 'Comunidade 1', 'Comunidade 2', 'Comunidade 3', 'Nome', 'Cargo', 'Organização',
             'LinkedIn', 'ID (não editar)']
    faltando = [c for c in obrig if c not in col]
    if faltando:
        raise ValueError(f'RNQ-selecao.xlsx: colunas ausentes na aba Pessoas: {faltando}')
    pessoas = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if txt(r[col['Incluir?']]).lower() != 'sim':
            continue
        nome = txt(r[col['Nome']])
        if not nome:
            continue
        pessoas.append({
            'nome': nome,
            'cargo': txt(r[col['Cargo']]),
            'organizacao': txt(r[col['Organização']]),
            'linkedin': txt(r[col['LinkedIn']]),
            'ext': txt(r[col['ID (não editar)']]),
            'comunidades': [txt(r[col[k]]) for k in ('Comunidade 1', 'Comunidade 2', 'Comunidade 3')
                            if txt(r[col[k]])],
        })
    ligacoes = []
    if 'Ligações' in wb.sheetnames:
        for r in wb['Ligações'].iter_rows(min_row=2, values_only=True):
            if r and len(r) >= 2 and txt(r[0]) and txt(r[1]):
                ligacoes.append((txt(r[0]), txt(r[1])))
    return {'pessoas': pessoas, 'comunidades': comunidades, 'ligacoes': ligacoes}


def ler_temas(caminho):
    """RNQ-temas-atuacao.xlsx -> {'por_chave': {ext ou nome: (temas, atuacao)}, 'ordem_temas', 'ordem_atuacao'}.
    Linha com "Aprovado? = Não" é ignorada (RF10)."""
    vazio = {'por_chave': {}, 'ordem_temas': [], 'ordem_atuacao': [], 'existe': False}
    if not caminho or not os.path.exists(caminho):
        return vazio
    wb = _abrir(caminho)
    ordem_temas = [txt(c.value) for c in wb['Temas']['A'][1:] if txt(c.value)]
    ordem_atu = [txt(c.value) for c in wb['Atuação']['A'][1:] if txt(c.value)]
    por_chave = {}
    for r in wb['Classificação'].iter_rows(min_row=2, values_only=True):
        if not r or not txt(r[0]) or nao(r[7]):
            continue
        temas = [txt(x) for x in r[1:4] if txt(x)]
        atu = [txt(x) for x in r[4:6] if txt(x)]
        if len(r) > 9 and txt(r[9]):
            por_chave[txt(r[9])] = (temas, atu)
        por_chave.setdefault(txt(r[0]), (temas, atu))
    return {'por_chave': por_chave, 'ordem_temas': ordem_temas, 'ordem_atuacao': ordem_atu, 'existe': True}


def ler_tecelagem(caminho):
    """RNQ-tecelagem.xlsx -> {aba: [ {coluna: valor}, ... ]}. Linhas totalmente vazias são ignoradas."""
    if not caminho or not os.path.exists(caminho):
        return {}
    wb = _abrir(caminho)
    out = {}
    for aba in wb.sheetnames:
        if aba == 'Como usar':
            continue
        ws = wb[aba]
        cab = [txt(c.value) for c in ws[1]]
        linhas = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or all(v is None or txt(v) == '' for v in r):
                continue
            linhas.append({cab[i]: r[i] for i in range(min(len(cab), len(r))) if cab[i]})
        out[aba] = linhas
    return out
