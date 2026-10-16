"""Planilha de tecelagem (RNQ-tecelagem.xlsx): onde chega tudo o que as pessoas enviam pelo site
e onde o Farley (e quem tecer junto) modera, registra apresentações e corrige dados.

Regras de desenho:
- Só validações de lista *inline* ("Sim,Não"): o Excel converte listas que apontam para outra aba
  num formato que o openpyxl apaga ao salvar — e esta planilha é reescrita pelo sincronizar.py.
- Colunas "(privado)" e "Observação interna" nunca vão para o site.
- O sincronizador só acrescenta linhas novas (deduplica pelo ID); nunca altera o que a pessoa editou.
"""
import glob
import os
import shutil
from datetime import datetime, timedelta, timezone

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .config import PAPEIS, STATUS_APRESENTACAO, TIPOS_INICIATIVA
from .util import txt

F = 'Arial'
COR_TOPO = '3B2A20'
COR_EDITAR = 'FFF2CC'
COR_PRIVADO = 'F4CCCC'
COR_INTERNO = 'EDEDED'

SN = '"Sim,Não"'
# aba -> (formulário do Netlify, [(coluna, largura, tipo, validação)])
#   tipo: 'auto' (preenchido pelo sincronizador), 'editar' (moderação), 'privado', 'interno'
ABAS = {
    'Mural': ('mural', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Tipo', 10, 'editar', '"Oferta,Pedido"'), ('Título', 34, 'editar', None),
        ('Descrição', 50, 'editar', None), ('Áreas', 34, 'editar', None),
        ('Nome', 24, 'editar', None), ('Contato (privado)', 26, 'privado', None),
        ('Está no mapa como', 24, 'editar', None), ('Publicar?', 10, 'editar', SN),
        ('Status', 11, 'editar', '"Aberto,Atendido"'), ('Observação interna', 30, 'interno', None)]),
    'Respostas do mural': ('mural-resposta', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('ID do item', 14, 'auto', None), ('Item', 30, 'auto', None), ('Nome', 24, 'auto', None),
        ('Contato (privado)', 26, 'privado', None), ('Mensagem', 44, 'auto', None),
        ('Apresentação feita?', 14, 'editar', SN), ('Observação interna', 30, 'interno', None)]),
    'Perfis': ('perfil', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Pessoa no mapa', 26, 'editar', None), ('Posso compartilhar', 36, 'editar', None),
        ('Quero aprender', 36, 'editar', None), ('Colaboro com', 30, 'editar', None),
        ('Gostaria de colaborar com', 30, 'interno', None), ('Toparia convocar grupo sobre', 30, 'editar', None),
        ('Contato (privado)', 24, 'privado', None), ('Publicar?', 10, 'editar', SN),
        ('Observação interna', 30, 'interno', None)]),
    'Ligações sugeridas': ('sugestao-ligacao', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Pessoa', 26, 'editar', None), ('Quem sugere', 26, 'editar', None),
        ('Contexto', 44, 'interno', None), ('Aprovar?', 10, 'editar', SN),
        ('Observação interna', 30, 'interno', None)]),
    'Saídas': ('opt-out', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Nome', 28, 'auto', None), ('Origem', 12, 'auto', None),
        ('Remover?', 10, 'editar', SN), ('Observação interna', 36, 'interno', None)]),
    'Confirmações': ('opt-in', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None), ('Nome', 28, 'auto', None)]),
    'Iniciativas': ('iniciativa', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Tipo', 16, 'editar', '"' + ','.join(TIPOS_INICIATIVA) + '"'), ('Título', 34, 'editar', None),
        ('Descrição', 50, 'editar', None), ('Áreas', 30, 'editar', None),
        ('Participantes', 34, 'editar', None), ('Precisa de', 34, 'editar', None),
        ('Nome de quem propôs', 24, 'editar', None), ('Contato (privado)', 24, 'privado', None),
        ('Publicar?', 10, 'editar', SN), ('Status', 12, 'editar', '"Ativa,Concluída"'),
        ('Observação interna', 30, 'interno', None)]),
    'Participações': ('iniciativa-participar', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('ID da iniciativa', 14, 'auto', None), ('Iniciativa', 30, 'auto', None), ('Nome', 24, 'auto', None),
        ('Contato (privado)', 26, 'privado', None), ('Mensagem', 44, 'auto', None),
        ('Encaminhado?', 12, 'editar', SN)]),
    'Cadastros': ('participar', [
        ('ID', 14, 'auto', None), ('Recebido em', 17, 'auto', None),
        ('Nome', 26, 'editar', None), ('Respondendo por', 14, 'auto', None),
        ('Cargo', 24, 'editar', None), ('Organização', 24, 'editar', None), ('LinkedIn', 26, 'editar', None),
        ('Site da organização', 24, 'editar', None), ('Tipo de organização', 18, 'editar', None),
        ('Contato (privado)', 26, 'privado', None), ('Cidade (privado)', 18, 'privado', None),
        ('Aniversário (privado)', 12, 'privado', None),
        ('Propósito', 44, 'editar', None), ('Temas', 40, 'editar', None), ('Tema novo sugerido', 22, 'interno', None),
        ('Atuação', 30, 'editar', None), ('Comunidades', 34, 'editar', None),
        ('Comunidade nova sugerida', 22, 'interno', None),
        ('Quer se aproximar de (privado)', 30, 'privado', None), ('Necessidade do território', 26, 'interno', None),
        ('Ofertas', 34, 'editar', None), ('Oferta (texto)', 40, 'editar', None), ('Pedido', 40, 'interno', None),
        ('Mural autorizado?', 10, 'auto', None), ('Objetivo', 16, 'interno', None), ('Fios', 30, 'editar', None),
        ('Formatos', 30, 'interno', None), ('Indicações (privado)', 40, 'privado', None),
        ('Visibilidade', 14, 'auto', None), ('18 anos ou mais', 10, 'auto', None), ('Idioma', 8, 'auto', None),
        ('Origem', 12, 'auto', None),
        ('Incluir no mapa?', 10, 'editar', SN), ('Observação interna', 30, 'interno', None)]),
    'Apresentações': (None, [
        ('Data', 12, 'editar', None), ('Pessoa A', 26, 'editar', None), ('Pessoa B', 26, 'editar', None),
        ('Motivo', 44, 'editar', None), ('Feita por', 20, 'editar', None),
        ('Status', 14, 'editar', '"' + ','.join(STATUS_APRESENTACAO) + '"'),
        ('Observação', 36, 'interno', None)]),
    'Papéis': (None, [
        ('Pessoa', 30, 'editar', None),
        ('Papel 1', 28, 'editar', '"' + ','.join(p['nome'] for p in PAPEIS) + '"'),
        ('Papel 2', 28, 'editar', '"' + ','.join(p['nome'] for p in PAPEIS) + '"')]),
    'Correções': (None, [
        ('Pessoa (nome na seleção)', 34, 'editar', None),
        ('Campo', 14, 'editar', '"Nome,Cargo,Organização,LinkedIn"'),
        ('Novo valor', 40, 'editar', None), ('Aplicar?', 10, 'editar', SN), ('Motivo', 40, 'interno', None)]),
}

# Formulário do Netlify -> (aba, {campo do formulário: coluna}, valores padrão)
FORMULARIOS = {
    'mural': ('Mural', {'tipo': 'Tipo', 'titulo': 'Título', 'descricao': 'Descrição', 'areas': 'Áreas',
                        'nome': 'Nome', 'contato': 'Contato (privado)', 'no_mapa': 'Está no mapa como'},
              {'Status': 'Aberto'}),
    'mural-resposta': ('Respostas do mural', {'item_id': 'ID do item', 'item_titulo': 'Item', 'nome': 'Nome',
                                              'contato': 'Contato (privado)', 'mensagem': 'Mensagem'}, {}),
    'perfil': ('Perfis', {'pessoa': 'Pessoa no mapa', 'compartilha': 'Posso compartilhar',
                          'aprende': 'Quero aprender', 'colabora_com': 'Colaboro com',
                          'gostaria_colaborar': 'Gostaria de colaborar com',
                          'convocaria': 'Toparia convocar grupo sobre', 'contato': 'Contato (privado)'}, {}),
    'sugestao-ligacao': ('Ligações sugeridas', {'pessoa': 'Pessoa', 'quem_sugere': 'Quem sugere',
                                                'contexto': 'Contexto'}, {}),
    # RF8: saída tem prioridade — entra marcada para remover; o Farley só desmarca se for abuso.
    'opt-out': ('Saídas', {'nome': 'Nome', 'origem': 'Origem'}, {'Remover?': 'Sim'}),
    'opt-in': ('Confirmações', {'nome': 'Nome'}, {}),
    'iniciativa': ('Iniciativas', {'tipo': 'Tipo', 'titulo': 'Título', 'descricao': 'Descrição',
                                   'areas': 'Áreas', 'participantes': 'Participantes', 'precisa': 'Precisa de',
                                   'nome': 'Nome de quem propôs', 'contato': 'Contato (privado)'},
                   {'Status': 'Ativa'}),
    'participar': ('Cadastros', {
        'nome': 'Nome', 'quem': 'Respondendo por', 'cargo': 'Cargo', 'organizacao': 'Organização',
        'linkedin': 'LinkedIn', 'org_site': 'Site da organização', 'org_tipo': 'Tipo de organização',
        'contato': 'Contato (privado)', 'local': 'Cidade (privado)', 'aniversario': 'Aniversário (privado)',
        'proposito': 'Propósito', 'temas': 'Temas', 'tema_outro': 'Tema novo sugerido', 'atuacao': 'Atuação',
        'comunidades': 'Comunidades', 'comunidade_outra': 'Comunidade nova sugerida',
        'quer_orgs': 'Quer se aproximar de (privado)', 'necessidade': 'Necessidade do território',
        'ofertas': 'Ofertas', 'oferta': 'Oferta (texto)', 'pedido': 'Pedido', 'mural': 'Mural autorizado?',
        'objetivo': 'Objetivo', 'fios': 'Fios', 'formatos': 'Formatos', 'indicacoes': 'Indicações (privado)',
        'visibilidade': 'Visibilidade', 'maior_18': '18 anos ou mais', 'idioma': 'Idioma', 'origem': 'Origem'},
        {}),
    'iniciativa-participar': ('Participações', {'iniciativa_id': 'ID da iniciativa', 'iniciativa_titulo': 'Iniciativa',
                                                'nome': 'Nome', 'contato': 'Contato (privado)',
                                                'mensagem': 'Mensagem'}, {}),
}

COMO_USAR = [
    ('Planilha de tecelagem da RNQ', True),
    ('', False),
    ('Tudo o que as pessoas enviam pelo site chega aqui (rode "python3 atualizar.py" ou peça ao Claude: "atualiza o mapa").', False),
    ('Nada vai para o site sem passar por você, exceto pedidos de saída do mapa (RF8: saem na próxima atualização).', False),
    ('', False),
    ('Cores: amarelo = você edita · vermelho = PRIVADO, nunca publicado · cinza = interno, nunca publicado.', True),
    ('', False),
    ('Mural ............ ofertas e pedidos. "Publicar?" = Sim publica; "Status" = Atendido tira o destaque.', False),
    ('Respostas do mural  alguém disse "posso ajudar". Faça a apresentação e marque "Apresentação feita?".', False),
    ('Perfis ........... respostas ao "Atualize seu perfil". "Publicar?" = Sim publica compartilha/aprende/convoca;', False),
    ('                   "Colaboro com" vira ligação no mapa; "Gostaria de colaborar com" alimenta o painel (privado).', False),
    ('Ligações sugeridas  alguém clicou "Eu conheço essa pessoa". "Aprovar?" = Sim cria a ligação.', False),
    ('Saídas ........... pedidos de saída. Já entram com "Remover?" = Sim. Só troque para Não se for abuso.', False),
    ('Confirmações ..... quem clicou "Sim, quero ficar no mapa".', False),
    ('Iniciativas ...... projetos, encontros e grupos de prática propostos. "Publicar?" = Sim publica.', False),
    ('Participações .... quem quer entrar numa iniciativa. Encaminhe para quem propôs e marque "Encaminhado?".', False),
    ('Cadastros ........ respostas ao "Faça parte". "Incluir no mapa?" = Sim põe a pessoa no mapa (se ela escolheu', False),
    ('                   aparecer e tem 18+). Oferta e pedido autorizados também chegam na aba Mural para revisar.', False),
    ('Apresentações .... seu diário de tecelagem. Registre cada apresentação e o status.', False),
    ('                   Apresentadas / Acompanhando / Colaborando = as duas pessoas passam a aparecer ligadas.', False),
    ('Papéis ........... selos públicos de papel na rede (Network Weaver Handbook). Só com o OK da pessoa.', False),
    ('Correções ........ corrige nome, cargo, organização ou LinkedIn sem mexer na planilha de seleção.', False),
    ('', False),
    ('Áreas e participantes: separe por ponto e vírgula (;). Use os nomes de temas/comunidades e de pessoas do mapa.', False),
]


class PlanilhaAberta(Exception):
    pass


def aberta_no_excel(caminho):
    """O Excel cria '~$<nome>' enquanto a planilha está aberta. Escrever nela nessa hora faz o Excel,
    ao salvar, apagar o que o sincronizador acrescentou."""
    pasta, nome = os.path.split(caminho)
    return os.path.exists(os.path.join(pasta, '~$' + nome))


def backup(caminho, manter=10):
    """Cópia de segurança em _trabalho/backups antes de cada gravação (guarda as 10 mais recentes)."""
    if not os.path.exists(caminho):
        return None
    pasta = os.path.join(os.path.dirname(caminho), '_trabalho', 'backups')
    os.makedirs(pasta, exist_ok=True)
    base = os.path.splitext(os.path.basename(caminho))[0]
    destino = os.path.join(pasta, f'{base}-{datetime.now().strftime("%Y%m%d-%H%M%S-%f")}.xlsx')
    shutil.copy2(caminho, destino)
    for velho in sorted(glob.glob(os.path.join(pasta, base + '-*.xlsx')))[:-manter]:
        try:
            os.remove(velho)
        except OSError:
            pass
    return destino


def _estilo_cabecalho(c):
    c.font = Font(name=F, bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=COR_TOPO)
    c.alignment = Alignment(wrap_text=True, vertical='center')


def _cor(tipo):
    return {'editar': COR_EDITAR, 'privado': COR_PRIVADO, 'interno': COR_INTERNO}.get(tipo)


def _preparar_aba(ws, colunas, linhas_validacao=1000):
    for j, (nome, largura, tipo, val) in enumerate(colunas, 1):
        c = ws.cell(1, j, nome)
        _estilo_cabecalho(c)
        ws.column_dimensions[get_column_letter(j)].width = largura
        if val:
            dv = DataValidation(type='list', formula1=val, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f'{get_column_letter(j)}2:{get_column_letter(j)}{linhas_validacao}')
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 30


def criar(caminho, papeis_iniciais=None, correcoes_iniciais=None):
    """Cria a planilha de tecelagem vazia (com papéis e correções iniciais, se houver)."""
    wb = openpyxl.Workbook()
    w0 = wb.active
    w0.title = 'Como usar'
    for i, (t, negrito) in enumerate(COMO_USAR, 1):
        w0.cell(i, 1, t).font = Font(name='Courier New' if '....' in t or t.startswith('   ') else F,
                                     bold=negrito, size=12 if i == 1 else 10)
    w0.column_dimensions['A'].width = 130
    for aba, (_, colunas) in ABAS.items():
        ws = wb.create_sheet(aba)
        _preparar_aba(ws, colunas)
    for linha in papeis_iniciais or []:
        _anexar(wb['Papéis'], ABAS['Papéis'][1], linha)
    for linha in correcoes_iniciais or []:
        _anexar(wb['Correções'], ABAS['Correções'][1], linha)
    wb.save(caminho)
    return caminho


def garantir(caminho, **kw):
    """Cria se não existir; se existir, acrescenta abas novas que ainda não estejam lá (migração)."""
    if not os.path.exists(caminho):
        return criar(caminho, **kw)
    wb = openpyxl.load_workbook(caminho)
    mudou = False
    for aba, (_, colunas) in ABAS.items():
        if aba not in wb.sheetnames:
            _preparar_aba(wb.create_sheet(aba), colunas)
            mudou = True
    if mudou:
        wb.save(caminho)
    return caminho


def _anexar(ws, colunas, valores):
    """valores: dict coluna->valor. Pinta as células conforme o tipo."""
    r = ws.max_row + 1
    if r == 2 and all(ws.cell(2, j).value is None for j in range(1, len(colunas) + 1)):
        r = 2
    for j, (nome, _, tipo, _) in enumerate(colunas, 1):
        c = ws.cell(r, j, valores.get(nome))
        c.font = Font(name=F, size=10)
        c.alignment = Alignment(wrap_text=tipo in ('editar', 'privado', 'interno') and nome not in ('ID', 'Publicar?'),
                                vertical='top')
        cor = _cor(tipo)
        if cor:
            c.fill = PatternFill('solid', fgColor=cor)
    return r


def ids_existentes(ws):
    col = None
    for j, c in enumerate(ws[1], 1):
        if txt(c.value) == 'ID':
            col = j
    if not col:
        return set()
    return {txt(ws.cell(r, col).value) for r in range(2, ws.max_row + 1) if txt(ws.cell(r, col).value)}


def adicionar_envios(caminho, envios):
    """envios: lista de dicts {'form': nome, 'id': str, 'criado': iso, 'dados': {campo: valor}}.
    Acrescenta só os que ainda não estão na planilha. Devolve {aba: quantidade_nova}."""
    if aberta_no_excel(caminho):
        raise PlanilhaAberta(f'Feche {os.path.basename(caminho)} no Excel antes de sincronizar '
                             '(senão, ao salvar, o Excel apaga as linhas novas).')
    garantir(caminho)
    wb = openpyxl.load_workbook(caminho)
    novos = {}
    vistos = {aba: ids_existentes(wb[aba]) for aba in wb.sheetnames if aba in ABAS}
    for e in sorted(envios, key=lambda x: x.get('criado') or ''):
        form = e.get('form')
        if form not in FORMULARIOS:
            continue
        aba, mapa, padrao = FORMULARIOS[form]
        eid = txt(e.get('id'))
        if not eid or eid in vistos[aba]:
            continue
        dados = e.get('dados') or {}
        if txt(dados.get('bot')):           # honeypot preenchido = robô
            continue
        valores = {'ID': eid, 'Recebido em': _quando(e.get('criado'))}
        valores.update(padrao)
        for campo, coluna in mapa.items():
            v = dados.get(campo)
            if isinstance(v, list):
                v = '; '.join(txt(x) for x in v)
            valores[coluna] = txt(v) or valores.get(coluna)
        if form == 'mural' and txt(valores.get('Tipo')).lower() in ('oferta', 'pedido'):
            valores['Tipo'] = txt(valores['Tipo']).capitalize()
        _anexar(wb[aba], ABAS[aba][1], valores)
        vistos[aba].add(eid)
        novos[aba] = novos.get(aba, 0) + 1
        if form == 'participar':
            for linha in mural_do_cadastro(eid, valores):
                if linha['ID'] not in vistos['Mural']:
                    _anexar(wb['Mural'], ABAS['Mural'][1], linha)
                    vistos['Mural'].add(linha['ID'])
                    novos['Mural'] = novos.get('Mural', 0) + 1
    if novos:
        backup(caminho)
        wb.save(caminho)
    return novos


def _titulo(texto, limite=90):
    t = txt(texto).split('\n')[0]
    for sep in ('. ', '; '):
        if sep in t:
            t = t.split(sep)[0]
    return t if len(t) <= limite else t[:limite - 1].rstrip() + '…'


def mural_do_cadastro(eid, c):
    """Oferta e pedido do "Faça parte" viram itens do Mural (sem "Publicar?": a moderação decide),
    só quando a pessoa autorizou o mural."""
    if txt(c.get('Mural autorizado?')).lower() not in ('sim', 'yes', 'true', 'on'):
        return []
    areas = '; '.join([a for a in [x.strip() for x in txt(c.get('Temas')).split(';')] if a][:3])
    no_mapa = c.get('Nome') if txt(c.get('Visibilidade')).lower() == 'mapa' else None
    base = {'Áreas': areas, 'Nome': c.get('Nome'), 'Contato (privado)': c.get('Contato (privado)'),
            'Está no mapa como': no_mapa, 'Status': 'Aberto', 'Recebido em': c.get('Recebido em'),
            'Observação interna': 'Veio do cadastro "Faça parte" (aba Cadastros)'}
    out = []
    ofertas, texto = txt(c.get('Ofertas')), txt(c.get('Oferta (texto)'))
    if ofertas or texto:
        titulo = _titulo(texto) if texto else 'Ofereço: ' + ofertas.replace(';', ',')
        out.append(dict(base, ID=eid + '-oferta', Tipo='Oferta', **{'Título': _titulo(titulo),
                   'Descrição': '; '.join(x for x in (texto, ofertas if texto else '') if x)}))
    if txt(c.get('Pedido')):
        out.append(dict(base, ID=eid + '-pedido', Tipo='Pedido', **{'Título': _titulo(c.get('Pedido')),
                   'Descrição': txt(c.get('Pedido'))}))
    return out


BRASILIA = timezone(timedelta(hours=-3))      # sem horário de verão desde 2019


def _quando(v):
    """Data do envio no horário de Brasília ('AAAA-MM-DD HH:MM'). O Netlify manda em UTC ('...Z');
    datas sem fuso (CSV, planilha) ficam como estão."""
    if not v:
        return datetime.now(BRASILIA).strftime('%Y-%m-%d %H:%M')
    s = txt(v)
    try:
        d = datetime.fromisoformat(s.replace('Z', '+00:00'))
        if d.tzinfo:
            return d.astimezone(BRASILIA).strftime('%Y-%m-%d %H:%M')
    except ValueError:
        pass
    return s.replace('T', ' ')[:16]
