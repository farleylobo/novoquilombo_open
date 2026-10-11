"""Ponta a ponta: planilhas de verdade (fictícias) -> arquivos do site e painel."""
import json
import os
import re
import shutil
import tempfile
import unittest
from datetime import datetime

import apoio

import exportar
from rnq import config as C
from rnq import tecelagem


class Exportar(unittest.TestCase):
    def setUp(self):
        self.d = tempfile.mkdtemp()
        apoio.escrever_selecao(os.path.join(self.d, 'RNQ-selecao.xlsx'))
        apoio.escrever_temas(os.path.join(self.d, 'RNQ-temas-atuacao.xlsx'))

    def tearDown(self):
        shutil.rmtree(self.d)

    def rodar(self, **kw):
        return exportar.executar(pasta=self.d, agora=datetime(2026, 10, 6, 9), silencioso=True, **kw)

    def test_gera_site_painel_historico_e_tecelagem(self):
        pub, avisos, interno, sug = self.rodar()
        for rel in ('site/rnq-publico.json', 'site/data.js', 'painel/painel-tecelagem.html',
                    '_trabalho/historico-rede.json', 'RNQ-tecelagem.xlsx'):
            self.assertTrue(os.path.exists(os.path.join(self.d, rel)), rel)
        js = apoio.ler(os.path.join(self.d, 'site/data.js'))
        self.assertTrue(js.startswith('window.RNQ_DATA = {'))
        self.assertEqual(json.loads(js[len('window.RNQ_DATA = '):].rstrip().rstrip(';')), pub)
        self.assertEqual(len(pub['pessoas']), 5)          # "Fora Da Seleção" não entra (RF2)
        self.assertNotIn('Fora Da Seleção', js)

    def test_ca1_nenhum_campo_interno_no_json(self):
        self.rodar()
        texto = apoio.ler(os.path.join(self.d, 'site/rnq-publico.json'))
        for proibido in ('Sinergia', 'Alta', 'interno', 'nota privada', 'base interna', 'Confiança',
                         'score', 'triage', 'attributes', 'email', 'connectedAt', 'li:ana-lima'):
            self.assertNotIn(proibido, texto, proibido)
        d = json.loads(texto)
        self.assertLessEqual(set(d), C.CAMPOS_TOPO)
        for p in d['pessoas']:
            self.assertLessEqual(set(p), C.CAMPOS_PESSOA)

    def test_fluxo_completo_de_moderacao(self):
        arq = os.path.join(self.d, 'RNQ-tecelagem.xlsx')
        tecelagem.criar(arq)
        tecelagem.adicionar_envios(arq, [
            {'form': 'mural', 'id': 'm-1', 'criado': '2026-10-05T10:00:00Z',
             'dados': {'tipo': 'pedido', 'titulo': 'Preciso de sala', 'descricao': 'Para 20 pessoas',
                       'areas': 'Clima & descarbonização', 'nome': 'Carla', 'no_mapa': 'Carla Dias',
                       'contato': 'carla-secreto@exemplo.org'}},
            {'form': 'opt-out', 'id': 'o-1', 'criado': '2026-10-05T11:00:00Z', 'dados': {'nome': 'Davi Nunes'}},
        ])
        pub, *_ = self.rodar()
        self.assertEqual(pub['mural'], [])                                  # ainda não aprovado
        self.assertNotIn('Davi Nunes', [p['nome'] for p in pub['pessoas']])  # saída já aplicada (RF8)
        import openpyxl
        wb = openpyxl.load_workbook(arq)
        ws = wb['Mural']
        ws.cell(2, [c.value for c in ws[1]].index('Publicar?') + 1, 'Sim')
        wb.save(arq)
        pub, *_ = self.rodar()
        self.assertEqual(pub['mural'][0]['titulo'], 'Preciso de sala')
        self.assertEqual(pub['mural'][0]['autor_id'], 'carla-dias')
        site = apoio.ler(os.path.join(self.d, 'site/rnq-publico.json'))
        self.assertNotIn('carla-secreto', site)
        painel = apoio.ler(os.path.join(self.d, 'painel/painel-tecelagem.html'))
        self.assertIn('carla-secreto@exemplo.org', painel)            # painel é privado e mostra o contato
        self.assertIn('Documento interno', painel)

    def test_historico_um_registro_por_dia(self):
        self.rodar()
        self.rodar()
        h = json.loads(apoio.ler(os.path.join(self.d, '_trabalho/historico-rede.json')))
        self.assertEqual(len(h), 1)
        self.assertEqual(h[0]['pessoas'], 5)

    def test_painel_nunca_sugere_o_hub(self):
        _, _, _, sug = self.rodar()
        self.assertFalse(any('farley-lobo' in (s['a'], s['b']) for s in sug))
        painel = apoio.ler(os.path.join(self.d, 'painel/painel-tecelagem.html'))
        self.assertIn('Apresentações sugeridas', painel)
        self.assertIsNone(re.search(r'<script[^>]+src=', painel), 'painel deve funcionar offline, sem scripts externos')


class SiteReal(unittest.TestCase):
    """Confere o JSON publicado de verdade (se existir), sem depender das planilhas."""

    def test_json_publicado_respeita_lista_fechada(self):
        arq = os.path.join(apoio.RAIZ, 'site', 'rnq-publico.json')
        if not os.path.exists(arq):
            self.skipTest('site ainda não exportado')
        d = json.loads(apoio.ler(arq))
        self.assertLessEqual(set(d), C.CAMPOS_TOPO)
        for p in d['pessoas']:
            self.assertLessEqual(set(p), C.CAMPOS_PESSOA, p.get('nome'))
        for m in d.get('mural', []):
            self.assertLessEqual(set(m), C.CAMPOS_MURAL)
        for i in d.get('iniciativas', []):
            self.assertLessEqual(set(i), C.CAMPOS_INICIATIVA)
        d.pop('sobre', None)          # único e-mail permitido: o contato oficial de privacidade (rnq/config.py)
        texto = json.dumps(d, ensure_ascii=False)
        self.assertIsNone(re.search(r'[\w.+-]+@[\w-]+\.[\w.]+', texto), 'há um e-mail no JSON público')


if __name__ == '__main__':
    unittest.main()
