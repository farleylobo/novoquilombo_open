"""Dados de teste pequenos e fictícios (nenhuma pessoa real)."""
import os
import sys

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, RAIZ)

import openpyxl  # noqa: E402


def ler(caminho):
    with open(caminho, encoding='utf-8') as fh:
        return fh.read()


def gravar(caminho, texto):
    with open(caminho, 'w', encoding='utf-8') as fh:
        fh.write(texto)


def pessoa(nome, coms=(), cargo='Cargo', org='Org', ext=None, linkedin=''):
    return {'nome': nome, 'cargo': cargo, 'organizacao': org, 'linkedin': linkedin,
            'ext': ext or 'li:' + nome.lower().replace(' ', '-'), 'comunidades': list(coms)}


def selecao_base():
    return {
        'comunidades': ['Cultura Regenerativa', 'IDGs', 'ESG', 'Tecnologia', 'Diversidade Equidade e Inclusao', 'YCL'],
        'pessoas': [
            pessoa('Ana Lima', ['Cultura Regenerativa', 'Diversidade Equidade e Inclusao'], org='Instituto Sol'),
            pessoa('Bruno Reis', ['IDGs'], org='Empresa X'),
            pessoa('Carla Dias', ['ESG'], org='Empresa X'),
            pessoa('Davi Nunes', ['YCL'], org='Coletivo Y'),
            pessoa('Farley Lobo', ['IDGs'], cargo='Fundador', org='Rede Novo Quilombo'),
        ],
        'ligacoes': [('Ana Lima', 'Farley Lobo'), ('Bruno Reis', 'Farley Lobo')],
    }


def temas_base():
    return {
        'existe': True,
        'ordem_temas': ['Clima & descarbonização', 'Produto, dados & transformação digital',
                        'Diversidade & equidade racial', 'Redes & ecossistemas'],
        'ordem_atuacao': ['Corporativo', 'Terceiro setor & coletivos'],
        'por_chave': {
            'li:ana-lima': (['Clima & descarbonização', 'Redes & ecossistemas'], ['Terceiro setor & coletivos']),
            'li:bruno-reis': (['Produto, dados & transformação digital', 'Clima & descarbonização'], ['Corporativo']),
            'li:carla-dias': (['Clima & descarbonização'], ['Corporativo']),
            'li:davi-nunes': (['Diversidade & equidade racial'], ['Terceiro setor & coletivos']),
            'li:farley-lobo': (['Redes & ecossistemas'], ['Terceiro setor & coletivos']),
        },
    }


def escrever_selecao(caminho, sel=None):
    sel = sel or selecao_base()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Pessoas'
    ws.append(['Incluir?', 'Comunidade 1', 'Comunidade 2', 'Comunidade 3', 'Nome', 'Cargo', 'Organização', 'Local',
               'Origem', 'Sinergia (interno)', 'Cluster Atria (interno)', 'LinkedIn', 'ID (não editar)'])
    for p in sel['pessoas']:
        c = list(p['comunidades']) + [None] * 3
        ws.append(['Sim', c[0], c[1], c[2], p['nome'], p['cargo'], p['organizacao'], 'SP', 'x', 'Alta', 'interno',
                   p['linkedin'], p['ext']])
    ws.append([None, None, None, None, 'Fora Da Seleção', 'x', 'y', '', '', '', '', '', 'li:fora'])
    wc = wb.create_sheet('Comunidades')
    wc.append(['Comunidade'])
    for c in sel['comunidades']:
        wc.append([c])
    wl = wb.create_sheet('Ligações')
    wl.append(['Pessoa A', 'Pessoa B', 'Observação (privada)'])
    for a, b in sel['ligacoes']:
        wl.append([a, b, 'nota privada da seleção'])
    wb.save(caminho)


def escrever_temas(caminho, tem=None, nomes=None):
    tem = tem or temas_base()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Classificação'
    ws.append(['Nome', 'Tema 1', 'Tema 2', 'Tema 3', 'Atuação 1', 'Atuação 2', 'Confiança', 'Aprovado?', 'Base', 'ID'])
    nomes = nomes or {'li:ana-lima': 'Ana Lima', 'li:bruno-reis': 'Bruno Reis', 'li:carla-dias': 'Carla Dias',
                      'li:davi-nunes': 'Davi Nunes', 'li:farley-lobo': 'Farley Lobo'}
    for ext, (t, a) in tem['por_chave'].items():
        t = list(t) + [None] * 3
        a = list(a) + [None] * 2
        ws.append([nomes[ext], t[0], t[1], t[2], a[0], a[1], 'Alta', None, 'base interna', ext])
    wt = wb.create_sheet('Temas')
    wt.append(['Temas'])
    for t in tem['ordem_temas']:
        wt.append([t])
    wa = wb.create_sheet('Atuação')
    wa.append(['Atuação'])
    for a in tem['ordem_atuacao']:
        wa.append([a])
    wb.save(caminho)
