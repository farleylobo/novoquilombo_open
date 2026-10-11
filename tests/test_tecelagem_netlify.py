import io
import json
import os
import tempfile
import unittest
import urllib.error

import apoio  # noqa: F401

import openpyxl

import sincronizar
from rnq import fontes, netlify, tecelagem


def envio(form, i, **dados):
    return {'form': form, 'id': i, 'criado': '2026-10-06T10:00:00Z', 'dados': dados}


class Planilha(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.arq = os.path.join(self.dir, 'RNQ-tecelagem.xlsx')

    def test_cria_todas_as_abas_com_cabecalhos(self):
        tecelagem.criar(self.arq, papeis_iniciais=[{'Pessoa': 'Ana', 'Papel 1': 'Conector(a)'}])
        wb = openpyxl.load_workbook(self.arq)
        for aba, (_, cols) in tecelagem.ABAS.items():
            self.assertIn(aba, wb.sheetnames)
            self.assertEqual([c.value for c in wb[aba][1]], [c[0] for c in cols])
        self.assertEqual(wb['Papéis']['A2'].value, 'Ana')

    def test_validacoes_sao_inline(self):
        # listas que apontam para outra aba se perdem quando o Excel salva e o openpyxl reescreve
        tecelagem.criar(self.arq)
        wb = openpyxl.load_workbook(self.arq)
        for ws in wb.worksheets:
            for dv in ws.data_validations.dataValidation:
                self.assertTrue(dv.formula1.startswith('"'), f'{ws.title}: {dv.formula1}')

    def test_adiciona_deduplica_e_preserva_edicoes(self):
        e1 = envio('mural', 'abc', tipo='oferta', titulo='Mentoria', nome='Ana', contato='ana@x.org',
                   areas=['Clima & descarbonização', 'ESG'])
        novos = tecelagem.adicionar_envios(self.arq, [e1])
        self.assertEqual(novos, {'Mural': 1})
        # Farley aprova
        wb = openpyxl.load_workbook(self.arq)
        ws = wb['Mural']
        col = [c.value for c in ws[1]].index('Publicar?') + 1
        ws.cell(2, col, 'Sim')
        wb.save(self.arq)
        # roda de novo com o mesmo envio + um novo
        e2 = envio('opt-out', 'def', nome='Bruno', origem='card')
        novos = tecelagem.adicionar_envios(self.arq, [e1, e2])
        self.assertEqual(novos, {'Saídas': 1})
        t = fontes.ler_tecelagem(self.arq)
        self.assertEqual(len(t['Mural']), 1)
        self.assertEqual(t['Mural'][0]['Publicar?'], 'Sim')
        self.assertEqual(t['Mural'][0]['Tipo'], 'Oferta')
        self.assertEqual(t['Mural'][0]['Áreas'], 'Clima & descarbonização; ESG')
        self.assertEqual(t['Mural'][0]['Status'], 'Aberto')
        self.assertEqual(t['Saídas'][0]['Remover?'], 'Sim')   # RF8

    def test_ignora_robo_e_formulario_desconhecido(self):
        novos = tecelagem.adicionar_envios(self.arq, [envio('mural', 'r1', titulo='spam', bot='x'),
                                                      envio('outro', 'z1', a='b')])
        self.assertEqual(novos, {})

    def test_garantir_migra_abas_novas(self):
        wb = openpyxl.Workbook()
        wb.active.title = 'Mural'
        wb.save(self.arq)
        tecelagem.garantir(self.arq)
        self.assertIn('Correções', openpyxl.load_workbook(self.arq).sheetnames)


class FakeResp(io.BytesIO):
    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False


def fake_api(respostas, erro=None):
    chamadas = []

    def abrir(req, timeout=30):
        chamadas.append(req.full_url)
        if erro:
            raise urllib.error.HTTPError(req.full_url, erro, 'x', {}, None)
        for chave, valor in respostas.items():
            if chave in req.full_url:
                return FakeResp(json.dumps(valor).encode())
        return FakeResp(b'[]')
    abrir.chamadas = chamadas
    return abrir


class Netlify(unittest.TestCase):
    def test_busca_todos_os_formularios_com_paginacao(self):
        pag1 = [{'id': f's{i}', 'created_at': '2026-10-06', 'data': {'nome': f'P{i}'}} for i in range(100)]
        abrir = fake_api({'/sites/SITE/forms': [{'id': 'F1', 'name': 'opt-in'}],
                          'F1/submissions?per_page=100&page=1': pag1,
                          'F1/submissions?per_page=100&page=2': [{'id': 'ult', 'data': {'nome': 'X'}}]})
        env = netlify.buscar_envios('TOKEN', 'SITE', abrir=abrir)
        self.assertEqual(len(env), 101)
        self.assertEqual(env[0]['form'], 'opt-in')

    def test_erros_claros(self):
        with self.assertRaises(netlify.ErroNetlify) as c:
            netlify.buscar_envios('T', 'S', abrir=fake_api({}, erro=401))
        self.assertIn('Token', str(c.exception))
        with self.assertRaises(netlify.ErroNetlify):
            netlify.buscar_envios('', '')

    def test_csv(self):
        d = tempfile.mkdtemp()
        arq = os.path.join(d, 'x.csv')
        apoio.gravar(arq, 'id,created_at,nome,origem\nab1,2026-10-01,Ana,card\n,2026-10-02,Bia,rodape\n')
        env = netlify.ler_csv(arq, 'opt-out')
        self.assertEqual(env[0]['id'], 'ab1')
        self.assertEqual(env[0]['dados']['nome'], 'Ana')
        self.assertTrue(env[1]['id'].startswith('csv-'))

    def test_env(self):
        d = tempfile.mkdtemp()
        apoio.gravar(os.path.join(d, '.env'), '# comentário\nNETLIFY_TOKEN="abc"\nNETLIFY_SITE_ID=xyz\n')
        cfg = netlify.ler_env(d)
        self.assertEqual(cfg['NETLIFY_TOKEN'], 'abc')
        self.assertEqual(cfg['NETLIFY_SITE_ID'], 'xyz')

    def test_sincronizar_ponta_a_ponta(self):
        d = tempfile.mkdtemp()
        apoio.gravar(os.path.join(d, '.env'), 'NETLIFY_TOKEN=t\nNETLIFY_SITE_ID=S\n')
        abrir = fake_api({'/sites/S/forms': [{'id': 'F', 'name': 'perfil'}],
                          'F/submissions?per_page=100&page=1': [{'id': 'p1', 'data': {'pessoa': 'Ana', 'aprende': 'Teatro'}}]})
        total, novos = sincronizar.executar(pasta=d, abrir=abrir)
        self.assertEqual((total, novos), (1, {'Perfis': 1}))
        self.assertEqual(fontes.ler_tecelagem(os.path.join(d, 'RNQ-tecelagem.xlsx'))['Perfis'][0]['Quero aprender'], 'Teatro')


class ContratoSiteFormularios(unittest.TestCase):
    """Os nomes dos campos nos formulários do site precisam bater com o que o sincronizador espera."""

    def test_campos_declarados_no_html(self):
        import re
        html = apoio.ler(os.path.join(apoio.RAIZ, 'site', 'index.html'))
        for form, (_, mapa, _) in tecelagem.FORMULARIOS.items():
            m = re.search(rf'<form name="{form}" data-netlify="true"[^>]*>(.*?)</form>', html, re.S)
            self.assertIsNotNone(m, f'formulário {form} não declarado no index.html')
            campos = set(re.findall(r'name="([a-z_0-9]+)"', m.group(1)))
            self.assertTrue(set(mapa) <= campos, f'{form}: faltam {set(mapa) - campos}')

    def test_campos_usados_no_app(self):
        js = apoio.ler(os.path.join(apoio.RAIZ, 'site', 'app.js')) + apoio.ler(os.path.join(apoio.RAIZ, 'site', 'participar.js'))
        for form, (_, mapa, _) in tecelagem.FORMULARIOS.items():
            self.assertIn(f"'{form}'", js, f'o site não envia o formulário {form}')

    def test_participar_envia_todos_os_campos_da_planilha(self):
        import re
        js = apoio.ler(os.path.join(apoio.RAIZ, 'site', 'participar.js'))
        bloco = re.search(r'function campos\(\)[^{]*\{(.*?)\n\}', js, re.S).group(1)
        enviados = set(re.findall(r'\b([a-z_0-9]+):', bloco))
        faltam = set(tecelagem.FORMULARIOS['participar'][1]) - enviados
        self.assertEqual(faltam, set(), f'participar.js não envia: {faltam}')


if __name__ == '__main__':
    unittest.main()
